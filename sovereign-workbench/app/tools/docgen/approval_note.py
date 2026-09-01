"""
tools/docgen/approval_note.py
==============================
Template-exact builder for Approval Note .docx documents.

This module is a pure builder: it takes structured data in, writes a Word
document out, and knows nothing about where the data came from (orchestrator,
RAG pipeline, manual input — irrelevant here).

Document structure (field-for-field match of the reference template)
---------------------------------------------------------------------
Heading 1   : APPROVAL NOTE  (centred)
Table        : Header block — Date / To / From / Reference No
Heading 2   : Section 1: Subject / Title
Paragraph   : subject_title  (blank placeholder if empty)
Heading 2   : Section 2: Background / Context
Paragraph   : background_context  (blank placeholder if empty)
Heading 2   : Section 3: Proposal / Request
Paragraph   : proposal_request  (blank placeholder if empty)
Heading 2   : Section 4: Justification / Business Case
Paragraph   : justification  (blank placeholder if empty)
Heading 2   : Approved By
Block        : Name/Designation: _______________
Heading 2   : Section 5: Financial Implications
Table        : cols: Budget Head / Cost Center | Estimated Cost
               plus a bold Total row at the bottom
Heading 2   : Initiated By
Block        : Name/Designation: _______________

Blank-field convention
----------------------
Any field with no data renders as "_______________" (15 underscores),
matching the template's own blank-field style.  Fields are never silently
omitted or filled with invented text.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

_BLANK = "_______________"  # 15 underscores — matches template blank convention

# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------


@dataclass
class FinancialRow:
    """One row in the Financial Implications table."""
    budget_head: str = ""
    estimated_cost: str = ""


@dataclass
class ApprovalNoteData:
    """
    All fields needed to produce a template-exact Approval Note .docx.

    Header block
    ------------
    date           : Date of the approval note
    to             : Addressee (name / designation)
    from_          : Originator (name / designation)
    reference_no   : Reference number of the approval note

    Sections
    --------
    subject_title       : Section 1 — free text
    background_context  : Section 2 — free text
    proposal_request    : Section 3 — free text
    justification       : Section 4 — free text (RAG citations go here)
    approved_by         : Approved By block — name/designation
    financial_rows      : Section 5 — list of FinancialRow
    financial_total     : Section 5 — total row value (string)
    initiated_by        : Initiated By block — name/designation
    """

    # Header block
    date: str = ""
    to: str = ""
    from_: str = ""
    reference_no: str = ""

    # Section 1
    subject_title: str = ""

    # Section 2
    background_context: str = ""

    # Section 3
    proposal_request: str = ""

    # Section 4
    justification: str = ""

    # Approved By (sits between sec 4 and sec 5)
    approved_by: str = ""

    # Section 5
    financial_rows: list[FinancialRow] = field(default_factory=list)
    financial_total: str = ""

    # Initiated By (end of document)
    initiated_by: str = ""


# ---------------------------------------------------------------------------
# Table column headers — exact match to reference template
# ---------------------------------------------------------------------------

_FIN_HEADERS = [
    "Budget Head / Cost Center",
    "Estimated Cost",
]


# ---------------------------------------------------------------------------
# Builder
# ---------------------------------------------------------------------------


def build_approval_note(data: ApprovalNoteData, path: Path) -> None:
    """
    Write a template-exact Approval Note .docx to ``path``.

    Parameters
    ----------
    data : ApprovalNoteData
        Structured content to populate the document with.
    path : Path
        Destination file path (parent directory must exist or will be created).
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # ── Document title ────────────────────────────────────────────────────
    title_para = doc.add_heading("APPROVAL NOTE", level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── Header block table ────────────────────────────────────────────────
    hdr_table = doc.add_table(rows=4, cols=2)
    hdr_table.style = "Table Grid"
    _fill_header_table(hdr_table, [
        ("Date", data.date or _BLANK),
        ("To", data.to or _BLANK),
        ("From", data.from_ or _BLANK),
        ("Reference No", data.reference_no or _BLANK),
    ])

    doc.add_paragraph("")  # spacer

    # ── Section 1: Subject / Title ────────────────────────────────────────
    doc.add_heading("Section 1: Subject / Title", level=2)
    doc.add_paragraph(data.subject_title if data.subject_title else _BLANK)

    # ── Section 2: Background / Context ──────────────────────────────────
    doc.add_heading("Section 2: Background / Context", level=2)
    doc.add_paragraph(data.background_context if data.background_context else _BLANK)

    # ── Section 3: Proposal / Request ────────────────────────────────────
    doc.add_heading("Section 3: Proposal / Request", level=2)
    doc.add_paragraph(data.proposal_request if data.proposal_request else _BLANK)

    # ── Section 4: Justification / Business Case ──────────────────────────
    doc.add_heading("Section 4: Justification / Business Case", level=2)
    doc.add_paragraph(data.justification if data.justification else _BLANK)

    doc.add_paragraph("")  # spacer

    # ── Approved By block ─────────────────────────────────────────────────
    doc.add_heading("Approved By", level=2)
    ab = data.approved_by if data.approved_by else _BLANK
    doc.add_paragraph(f"Name/Designation: {ab}")

    doc.add_paragraph("")  # spacer

    # ── Section 5: Financial Implications ────────────────────────────────
    doc.add_heading("Section 5: Financial Implications", level=2)
    fin_table = doc.add_table(rows=1, cols=2)
    fin_table.style = "Table Grid"
    _write_bold_header_row(fin_table.rows[0], _FIN_HEADERS)

    if data.financial_rows:
        for row_data in data.financial_rows:
            row = fin_table.add_row()
            row.cells[0].text = str(row_data.budget_head) if row_data.budget_head else _BLANK
            row.cells[1].text = str(row_data.estimated_cost) if row_data.estimated_cost else _BLANK
    else:
        row = fin_table.add_row()
        row.cells[0].text = _BLANK
        row.cells[1].text = _BLANK

    # Total row — always present, bold
    total_row = fin_table.add_row()
    total_cell_label = total_row.cells[0]
    total_cell_value = total_row.cells[1]
    _set_bold_cell_text(total_cell_label, "Total")
    _set_bold_cell_text(total_cell_value, data.financial_total if data.financial_total else _BLANK)

    doc.add_paragraph("")  # spacer

    # ── Initiated By block ────────────────────────────────────────────────
    doc.add_heading("Initiated By", level=2)
    ib = data.initiated_by if data.initiated_by else _BLANK
    doc.add_paragraph(f"Name/Designation: {ib}")

    doc.save(str(path))


# ---------------------------------------------------------------------------
# Private helpers
# ---------------------------------------------------------------------------


def _fill_header_table(table, rows: list[tuple[str, str]]) -> None:
    """Populate a 2-col label/value header table."""
    for i, (label, value) in enumerate(rows):
        cells = table.rows[i].cells
        # Set label with bold
        para = cells[0].paragraphs[0]
        para.clear()
        run = para.add_run(label)
        run.bold = True
        # Set value
        cells[1].text = value


def _write_bold_header_row(row, headers: list[str]) -> None:
    """Write bold column headers into a table row."""
    for i, header in enumerate(headers):
        _set_bold_cell_text(row.cells[i], header)


def _set_bold_cell_text(cell, text: str) -> None:
    """Clear a cell and write bold text into it."""
    para = cell.paragraphs[0]
    para.clear()
    run = para.add_run(text)
    run.bold = True
