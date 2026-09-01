"""
tools/docgen/inspection_report.py
==================================
Template-exact builder for Inspection Report .docx documents.

This module is a pure builder: it takes structured data in, writes a Word
document out, and knows nothing about where the data came from (orchestrator,
vision pipeline, manual input — irrelevant here).

Document structure (field-for-field match of the reference template)
---------------------------------------------------------------------
Heading 1   : INSPECTION REPORT  (centred)
Table        : Header block — Report No / Date of Inspection /
               Facility-Location / Inspector Name(s) / Equipment-Area
Heading 2   : Section 1: Objective of Inspection
Paragraph   : objective_text  (blank placeholder if empty)
Heading 2   : Section 2: Observations & Findings
Table        : cols: Sl No. | Checklist Item / Observation | Status (Pass/Fail) | Remarks
Heading 2   : Facility Manager Acknowledgment          ← mid-document, NOT at end
Block        : Signature: _______________ / Date: _______________
Heading 2   : Section 3: Non-Conformances (NCs) & Critical Issues
Paragraph   : non_conformances_text  (blank placeholder if empty)
Heading 2   : Section 4: Corrective Action / Recommendations
Table        : cols: Recommended Action | Responsibility | Target Date
Heading 2   : Lead Inspector
Block        : Signature: _______________ / Date: _______________

Blank-field convention
----------------------
Any field with no data renders as "_______________" (15 underscores),
matching the template's own blank-field style.  Fields are never silently
omitted or filled with invented text.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

_BLANK = "_______________"  # 15 underscores — matches template blank convention

# ---------------------------------------------------------------------------
# Data classes — callers fill these, builder reads them
# ---------------------------------------------------------------------------


@dataclass
class ObservationRow:
    """One row in the Observations & Findings table."""
    sl_no: str = ""
    checklist_item: str = ""
    status: str = ""        # Pass / Fail / N/A
    remarks: str = ""


@dataclass
class CorrectiveActionRow:
    """One row in the Corrective Action / Recommendations table."""
    recommended_action: str = ""
    responsibility: str = ""
    target_date: str = ""


@dataclass
class InspectionReportData:
    """
    All fields needed to produce a template-exact Inspection Report .docx.

    Header block
    ------------
    report_no            : Report reference number
    date_of_inspection   : ISO date string, e.g. "2026-08-15"
    facility_location    : Facility / Location name
    inspector_names      : Inspector name(s), comma-separated if multiple
    equipment_area       : Equipment or area inspected

    Sections
    --------
    objective            : Free text for Section 1
    observations         : List of ObservationRow for Section 2 table
    facility_manager_sig : Facility manager signature name (blank → placeholder)
    facility_manager_date: Facility manager signature date (blank → placeholder)
    non_conformances     : Free text for Section 3
    corrective_actions   : List of CorrectiveActionRow for Section 4 table
    lead_inspector_sig   : Lead inspector signature name (blank → placeholder)
    lead_inspector_date  : Lead inspector signature date (blank → placeholder)
    """

    # Header block
    report_no: str = ""
    date_of_inspection: str = ""
    facility_location: str = ""
    inspector_names: str = ""
    equipment_area: str = ""

    # Section 1
    objective: str = ""

    # Section 2
    observations: list[ObservationRow] = field(default_factory=list)

    # Facility Manager acknowledgment (between sec 2 and sec 3)
    facility_manager_sig: str = ""
    facility_manager_date: str = ""

    # Section 3
    non_conformances: str = ""

    # Section 4
    corrective_actions: list[CorrectiveActionRow] = field(default_factory=list)

    # Lead Inspector block (end of document)
    lead_inspector_sig: str = ""
    lead_inspector_date: str = ""


# ---------------------------------------------------------------------------
# Table column headers — exact match to reference template
# ---------------------------------------------------------------------------

_OBS_HEADERS = [
    "Sl No.",
    "Checklist Item / Observation",
    "Status (Pass/Fail)",
    "Remarks",
]

_CA_HEADERS = [
    "Recommended Action",
    "Responsibility",
    "Target Date",
]


# ---------------------------------------------------------------------------
# Builder
# ---------------------------------------------------------------------------


def build_inspection_report(data: InspectionReportData, path: Path) -> None:
    """
    Write a template-exact Inspection Report .docx to ``path``.

    Parameters
    ----------
    data : InspectionReportData
        Structured content to populate the document with.
    path : Path
        Destination file path (parent directory must exist or will be created).
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # ── Document title ────────────────────────────────────────────────────
    title_para = doc.add_heading("INSPECTION REPORT", level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── Header block table ────────────────────────────────────────────────
    hdr_table = doc.add_table(rows=5, cols=2)
    hdr_table.style = "Table Grid"
    _fill_header_table(hdr_table, [
        ("Report No", data.report_no or _BLANK),
        ("Date of Inspection", data.date_of_inspection or _BLANK),
        ("Facility / Location", data.facility_location or _BLANK),
        ("Inspector Name(s)", data.inspector_names or _BLANK),
        ("Equipment / Area", data.equipment_area or _BLANK),
    ])

    doc.add_paragraph("")  # spacer

    # ── Section 1: Objective of Inspection ───────────────────────────────
    doc.add_heading("Section 1: Objective of Inspection", level=2)
    doc.add_paragraph(data.objective if data.objective else _BLANK)

    # ── Section 2: Observations & Findings ───────────────────────────────
    doc.add_heading("Section 2: Observations & Findings", level=2)
    obs_table = doc.add_table(rows=1, cols=4)
    obs_table.style = "Table Grid"
    _write_bold_header_row(obs_table.rows[0], _OBS_HEADERS)

    if data.observations:
        for row_data in data.observations:
            row = obs_table.add_row()
            row.cells[0].text = str(row_data.sl_no) if row_data.sl_no else _BLANK
            row.cells[1].text = str(row_data.checklist_item) if row_data.checklist_item else _BLANK
            row.cells[2].text = str(row_data.status) if row_data.status else _BLANK
            row.cells[3].text = str(row_data.remarks) if row_data.remarks else _BLANK
    else:
        # At least one placeholder row so the table isn't completely empty
        row = obs_table.add_row()
        for cell in row.cells:
            cell.text = _BLANK

    doc.add_paragraph("")  # spacer

    # ── Facility Manager Acknowledgment (mid-document) ────────────────────
    doc.add_heading("Facility Manager Acknowledgment", level=2)
    sig = data.facility_manager_sig if data.facility_manager_sig else _BLANK
    sig_date = data.facility_manager_date if data.facility_manager_date else _BLANK
    doc.add_paragraph(f"Signature: {sig}")
    doc.add_paragraph(f"Date: {sig_date}")

    doc.add_paragraph("")  # spacer

    # ── Section 3: Non-Conformances (NCs) & Critical Issues ──────────────
    doc.add_heading("Section 3: Non-Conformances (NCs) & Critical Issues", level=2)
    doc.add_paragraph(data.non_conformances if data.non_conformances else _BLANK)

    # ── Section 4: Corrective Action / Recommendations ───────────────────
    doc.add_heading("Section 4: Corrective Action / Recommendations", level=2)
    ca_table = doc.add_table(rows=1, cols=3)
    ca_table.style = "Table Grid"
    _write_bold_header_row(ca_table.rows[0], _CA_HEADERS)

    if data.corrective_actions:
        for row_data in data.corrective_actions:
            row = ca_table.add_row()
            row.cells[0].text = str(row_data.recommended_action) if row_data.recommended_action else _BLANK
            row.cells[1].text = str(row_data.responsibility) if row_data.responsibility else _BLANK
            row.cells[2].text = str(row_data.target_date) if row_data.target_date else _BLANK
    else:
        row = ca_table.add_row()
        for cell in row.cells:
            cell.text = _BLANK

    doc.add_paragraph("")  # spacer

    # ── Lead Inspector signature block ────────────────────────────────────
    doc.add_heading("Lead Inspector", level=2)
    li_sig = data.lead_inspector_sig if data.lead_inspector_sig else _BLANK
    li_date = data.lead_inspector_date if data.lead_inspector_date else _BLANK
    doc.add_paragraph(f"Signature: {li_sig}")
    doc.add_paragraph(f"Date: {li_date}")

    doc.save(str(path))


# ---------------------------------------------------------------------------
# Private helpers
# ---------------------------------------------------------------------------


def _fill_header_table(table, rows: list[tuple[str, str]]) -> None:
    """Populate a 2-col label/value header table."""
    for i, (label, value) in enumerate(rows):
        cells = table.rows[i].cells
        cells[0].text = label
        # Bold the label
        run = cells[0].paragraphs[0].runs
        if run:
            run[0].bold = True
        cells[1].text = value


def _write_bold_header_row(row, headers: list[str]) -> None:
    """Write bold column headers into the first row of a table."""
    for i, header in enumerate(headers):
        cell = row.cells[i]
        cell.text = header
        runs = cell.paragraphs[0].runs
        if runs:
            runs[0].bold = True
        else:
            cell.paragraphs[0].add_run(header).bold = True
            cell.text = ""  # clear the text set above since we're using run
            cell.paragraphs[0].runs[0].text = header
            cell.paragraphs[0].runs[0].bold = True
