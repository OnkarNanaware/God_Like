"""
tests/test_approval_note_from_excel.py
=======================================
Regression tests for the Excel budget → approval note pipeline fix.

Bug A: analyze_spreadsheet was silently discarding actual row data,
       only passing aggregated stats (min/max/mean) to downstream steps.
Bug B: generate_docx.financial_rows was always empty because it was
       baked at plan-time before the file was read.
"""
from __future__ import annotations

import csv
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_budget_xlsx(rows: list) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    for row in rows:
        ws.append(row)
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tmp.name)
    tmp.close()
    return Path(tmp.name)

_SAMPLE_BUDGET_ROWS = [
    ("Budget Head / Cost Center", "Estimated Cost"),
    ("Marketing & Advertising", 45000),
    ("Human Resources",          62000),
    ("IT Infrastructure",        38500),
    ("Operations & Logistics",   27000),
    ("Research & Development",   55000),
]
_EXPECTED_TOTAL = 45000 + 62000 + 38500 + 27000 + 55000  # 227500

# ---------------------------------------------------------------------------
# Bug A: analyze_spreadsheet must include actual row data
# ---------------------------------------------------------------------------

class TestAnalyzeSpreadsheetRowData:

    def setup_method(self):
        from app.tools.analyze_spreadsheet import AnalyzeSpreadsheetTool
        fake_llm = MagicMock()
        fake_llm.chat_completion = AsyncMock(
            return_value=MagicMock(content="Budget spreadsheet with 5 line items.")
        )
        self.tool = AnalyzeSpreadsheetTool(llm_client=fake_llm)
        self.xlsx_path = _make_budget_xlsx(_SAMPLE_BUDGET_ROWS)

    def teardown_method(self):
        self.xlsx_path.unlink(missing_ok=True)

    @pytest.mark.asyncio
    async def test_rows_key_present(self):
        result = await self.tool.execute(file_path=str(self.xlsx_path))
        assert result.success, f"Tool failed: {result.error}"
        assert "rows" in result.output["sheets"][0], "Bug A: 'rows' key missing"

    @pytest.mark.asyncio
    async def test_rows_contain_line_items(self):
        result = await self.tool.execute(file_path=str(self.xlsx_path))
        rows = result.output["sheets"][0]["rows"]
        assert len(rows) == 5, f"Expected 5 data rows, got {len(rows)}"

    @pytest.mark.asyncio
    async def test_rows_have_correct_values(self):
        result = await self.tool.execute(file_path=str(self.xlsx_path))
        rows = result.output["sheets"][0]["rows"]
        first = rows[0]
        assert first["Budget Head / Cost Center"] == "Marketing & Advertising"
        assert first["Estimated Cost"] == 45000

    @pytest.mark.asyncio
    async def test_numeric_stats_still_present(self):
        result = await self.tool.execute(file_path=str(self.xlsx_path))
        stats = result.output["sheets"][0]["numeric_stats"]
        assert "Estimated Cost" in stats
        assert stats["Estimated Cost"]["min"] == 27000
        assert stats["Estimated Cost"]["max"] == 62000

# ---------------------------------------------------------------------------
# Bug B: generate_docx must hydrate financial rows from spreadsheet
# ---------------------------------------------------------------------------

class TestGenerateDocxFromSpreadsheet:

    def setup_method(self):
        self.xlsx_path = _make_budget_xlsx(_SAMPLE_BUDGET_ROWS)

    def teardown_method(self):
        self.xlsx_path.unlink(missing_ok=True)

    def test_rows_from_spreadsheet_returns_correct_rows(self):
        from app.tools.docgen.generate_docx import GenerateDocxTool
        rows, total_str = GenerateDocxTool._rows_from_spreadsheet(str(self.xlsx_path))
        assert len(rows) == 5, f"Expected 5 rows, got {len(rows)}"
        assert rows[0].budget_head == "Marketing & Advertising"
        assert "45" in rows[0].estimated_cost

    def test_rows_from_spreadsheet_computes_total(self):
        from app.tools.docgen.generate_docx import GenerateDocxTool
        rows, total_str = GenerateDocxTool._rows_from_spreadsheet(str(self.xlsx_path))
        assert total_str, "Total must not be empty"
        assert float(total_str.replace(",", "")) == pytest.approx(_EXPECTED_TOTAL)

    @pytest.mark.asyncio
    async def test_build_approval_note_hydrates_from_spreadsheet(self):
        from app.tools.docgen.generate_docx import GenerateDocxTool
        from docx import Document
        tool = GenerateDocxTool()
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = Path(tmpdir) / "approval_note.docx"
            tool._build_approval_note(
                path=out_path,
                kwargs={
                    "source_spreadsheet_path": str(self.xlsx_path),
                    "approval_data": {
                        "date": "2026-09-02",
                        "subject_title": "Annual Budget Approval",
                        "financial_rows": [],
                        "financial_total": "",
                    },
                },
            )
            assert out_path.exists()
            doc = Document(str(out_path))
            fin_table = doc.tables[1]
            data_cells = [row.cells[0].text for row in fin_table.rows[1:-1]]
            assert "Marketing & Advertising" in data_cells, (
                f"Bug B not fixed: financial rows not hydrated. Cells: {data_cells}"
            )
            total_val = fin_table.rows[-1].cells[1].text
            assert total_val and total_val != "_______________", (
                f"Total row still blank: {total_val!r}"
            )

    def test_empty_rows_are_skipped(self):
        rows_with_blanks = [
            ("Budget Head / Cost Center", "Estimated Cost"),
            ("Marketing", 45000),
            (None, None),
            ("", ""),
            ("IT", 38500),
        ]
        path = _make_budget_xlsx(rows_with_blanks)
        try:
            from app.tools.docgen.generate_docx import GenerateDocxTool
            rows, _ = GenerateDocxTool._rows_from_spreadsheet(str(path))
            assert len(rows) == 2, f"Expected 2, got {len(rows)}: {[r.budget_head for r in rows]}"
        finally:
            path.unlink(missing_ok=True)

    def test_csv_support(self):
        with tempfile.NamedTemporaryFile(
            suffix=".csv", mode="w", newline="", delete=False, encoding="utf-8"
        ) as fh:
            writer = csv.writer(fh)
            writer.writerow(["Department", "Budget"])
            writer.writerow(["HR", "62000"])
            writer.writerow(["IT", "38500"])
            csv_path = Path(fh.name)
        try:
            from app.tools.docgen.generate_docx import GenerateDocxTool
            rows, total_str = GenerateDocxTool._rows_from_spreadsheet(str(csv_path))
            assert len(rows) == 2
            assert rows[0].budget_head == "HR"
            assert float(total_str.replace(",", "")) == pytest.approx(100500.0)
        finally:
            csv_path.unlink(missing_ok=True)
