"""
tests/test_phase_d.py
=====================
Phase D test suite — multimodal input and real deliverable output.

Coverage
--------
 1.  VisionExtractTool: parses valid structured JSON from model
 2.  VisionExtractTool: retries and fails on persistent non-JSON output
 3.  VisionExtractTool: correct image-call count for a 3-page (mocked) PDF
 4.  VisionExtractTool: router fires vision path on PDF attachment
 5.  GenerateDocxTool: produces valid .docx with expected sections
 6.  GenerateDocxTool: findings table rows present in generated file
 7.  GenerateDocxTool: returns failure on missing required args
 8.  GeneratePptxTool: produces valid .pptx with expected slide count
 9.  GeneratePptxTool: slide headings match input
10.  GenerateXlsxTool: produces valid .xlsx with correct cell values
11.  GenerateXlsxTool: formula row written to sheet
12.  Orchestrator: generate_docx callable in a plan
13.  Orchestrator: generate_pptx callable in a plan
14.  Orchestrator: code_sandbox callable in a plan
15.  [Docker] CodeSandboxTool: trivial print script succeeds
16.  [Docker] CodeSandboxTool: exception-raising script returns exit 1
17.  [Docker] CodeSandboxTool: slow script times out
18.  [Docker] CodeSandboxTool: no leftover containers after run
"""

from __future__ import annotations

import asyncio
import json
import subprocess
import tempfile
import uuid
from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _run(coro):
    """Run a coroutine synchronously (for tests that cannot use pytest-asyncio)."""
    return asyncio.run(coro)


def _make_audit(tmp_path: Path):
    from app.audit.logger import AuditLogger
    return AuditLogger(log_path=tmp_path / "audit.jsonl")


# ---------------------------------------------------------------------------
# Fake vision client
# ---------------------------------------------------------------------------

class FakeVisionResponse:
    def __init__(self, content: str):
        self.content = content
        self.prompt_tokens = 10
        self.response_tokens = 20
        self.total_tokens = 30
        self.latency_ms = 1.0
        self.model_name = "fake-vision"
        self.ollama_tag = "fake:latest"
        self.request_id = "fake-rid"


class FakeVisionClient:
    """
    Returns canned JSON responses, cycling through a list.
    If a response is None it simulates a model exception.
    """
    def __init__(self, responses: list):
        self._responses = responses
        self._idx = 0
        self.call_count = 0

    async def chat_completion(self, messages, *, request_id=None,
                               temperature=0.0, max_tokens=2048,
                               extra_body=None):
        resp = self._responses[min(self._idx, len(self._responses) - 1)]
        self._idx += 1
        self.call_count += 1
        if resp is None:
            raise RuntimeError("Simulated model failure")
        return FakeVisionResponse(content=resp)


_GOOD_JSON = json.dumps({
    "page": 1,
    "fields": [
        {"name": "Report ID", "value": "MRPL-047", "region": "header"},
        {"name": "Date", "value": "2026-08-15", "region": "header"},
    ],
    "tables": [
        {
            "title": "Findings",
            "headers": ["ID", "Severity"],
            "rows": [["F-001", "HIGH"]],
        }
    ],
    "annotations": [],
})


# ===========================================================================
# 1–4: VisionExtractTool
# ===========================================================================

class TestVisionExtractTool:

    def _make_tool(self, responses, audit=None):
        from app.tools.vision_extract import VisionExtractTool
        client = FakeVisionClient(responses)
        return VisionExtractTool(llm_client=client, audit_logger=audit), client

    # ── Test 1: valid JSON response parsed correctly ─────────────────────

    def test_valid_json_response_parsed(self, tmp_path):
        """Tool returns success with correct field/table counts."""
        tool, client = self._make_tool([_GOOD_JSON])

        # Create a dummy PNG so the file-exists check passes
        img_path = tmp_path / "test.png"
        from PIL import Image
        Image.new("RGB", (100, 100), color=(255, 255, 255)).save(str(img_path))

        result = _run(tool.execute(file_path=str(img_path), request_id="test-001"))

        assert result.success is True, f"Expected success, got error: {result.error}"
        assert isinstance(result.output, list)
        assert len(result.output) == 1
        page = result.output[0]
        assert len(page["fields"]) == 2
        assert len(page["tables"]) == 1
        assert result.metadata["field_count"] == 2
        assert result.metadata["table_count"] == 1

    # ── Test 2: persistent non-JSON output → failure ─────────────────────

    def test_non_json_output_fails(self, tmp_path):
        """Tool exhausts retries and returns failure if model never returns JSON."""
        bad_responses = ["This is just prose text."] * 10  # more than MAX_RETRIES
        tool, client = self._make_tool(bad_responses)

        img_path = tmp_path / "test.png"
        from PIL import Image
        Image.new("RGB", (100, 100)).save(str(img_path))

        result = _run(tool.execute(file_path=str(img_path), request_id="test-002"))

        assert result.success is False
        assert "non-JSON" in result.error or "prose" in result.error.lower()

    # ── Test 3: multi-page PDF → correct call count ──────────────────────

    def test_pdf_rasterises_to_page_count(self, tmp_path):
        """
        A 3-page PDF must result in exactly 3 vision model calls.
        We mock convert_from_path to return 3 fake PIL images
        so we don't need poppler installed.
        """
        from PIL import Image
        fake_images = [
            Image.new("RGB", (100, 100), color=(i * 50, 0, 0))
            for i in range(3)
        ]

        # Each page gets _GOOD_JSON with correct page number
        responses = [
            json.dumps({**json.loads(_GOOD_JSON), "page": i + 1})
            for i in range(3)
        ]
        tool, client = self._make_tool(responses)

        fake_pdf = tmp_path / "report.pdf"
        fake_pdf.write_bytes(b"%PDF-1.4 fake")  # not real, but file exists

        with patch("pdf2image.convert_from_path", return_value=fake_images):
            result = _run(tool.execute(file_path=str(fake_pdf), request_id="test-003"))

        assert result.success is True, f"error: {result.error}"
        assert client.call_count == 3, f"Expected 3 calls, got {client.call_count}"
        assert result.metadata["page_count"] == 3

    # ── Test 4: router fires vision path on PDF attachment ───────────────

    def test_router_fires_vision_for_pdf_attachment(self):
        """
        run_heuristics() with a .pdf attachment returns a signal with
        confidence >= CONFIDENCE_THRESHOLD and a vision-related capability.
        """
        from app.router.heuristics import run_heuristics, is_ambiguous, Capability

        signals = run_heuristics(
            "Please extract the findings from this inspection document.",
            attached_filenames=["inspection_report.pdf"],
        )

        assert signals, "Expected at least one heuristic signal"
        assert not is_ambiguous(signals), (
            f"Signal was unexpectedly ambiguous (confidence={signals[0].confidence})"
        )
        top = signals[0]
        # PDF should map to DOCUMENT_VISION (or IMAGE_UNDERSTANDING for images)
        assert top.capability in (
            Capability.DOCUMENT_VISION,
            Capability.IMAGE_UNDERSTANDING,
            Capability.OCR,
        ), f"Unexpected capability: {top.capability}"
        assert top.confidence >= 0.75


# ===========================================================================
# 5–7: GenerateDocxTool
# ===========================================================================

class TestGenerateDocxTool:

    def _make_tool(self, tmp_path):
        from app.tools.docgen.generate_docx import GenerateDocxTool
        audit = _make_audit(tmp_path)
        return GenerateDocxTool(audit_logger=audit)

    # ── Test 5: valid .docx with expected sections ────────────────────────

    def test_generates_valid_docx_with_sections(self, tmp_path):
        """Tool produces a file that python-docx can open with expected headings."""
        tool = self._make_tool(tmp_path)

        result = _run(tool.execute(
            title="Test Approval Note",
            sections=[
                {"heading": "Background", "body": "Some background text."},
                {"heading": "Analysis", "body": "Detailed analysis here."},
            ],
            output_filename="test_output.docx",
            request_id="docx-001",
        ))

        assert result.success is True, f"Expected success, got: {result.error}"
        # result.output is now the sanitized filename; physical path is in the artifact
        assert "artifact" in result.metadata, "metadata['artifact'] must be present"
        art = result.metadata["artifact"]
        assert "physical_path" not in art, "physical_path must NOT appear in to_dict()"
        assert art["artifact_id"], "artifact_id must be non-empty"
        assert art["filename"].endswith(".docx")
        assert art["download_url"] == f"/outputs/{art['artifact_id']}"

        # Get the physical path directly from the ArtifactManager for content checks
        from app.artifacts.manager import get_artifact_manager
        artifact_obj = get_artifact_manager().get_artifact(art["artifact_id"])
        output_path = artifact_obj.physical_path
        assert output_path.exists(), "Output file not found"
        assert output_path.suffix == ".docx"

        # Re-open and verify structure
        from docx import Document
        doc = Document(str(output_path))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "Test Approval Note" in headings, f"Title missing from headings: {headings}"
        assert "Background" in headings, f"Section 'Background' missing: {headings}"
        assert "Analysis" in headings, f"Section 'Analysis' missing: {headings}"

    # ── Test 6: findings table present in generated file ─────────────────

    def test_generates_docx_with_findings_table(self, tmp_path):
        """Findings table rows are present and correctly populated."""
        tool = self._make_tool(tmp_path)

        result = _run(tool.execute(
            title="Findings Report",
            sections=[{"heading": "Overview", "body": "Overview of findings."}],
            findings_table=[
                {"id": "F-001", "finding": "Temperature anomaly", "severity": "HIGH", "status": "OPEN"},
                {"id": "F-002", "finding": "Pressure variance", "severity": "MEDIUM", "status": "OPEN"},
            ],
            output_filename="test_findings.docx",
            request_id="docx-002",
        ))

        assert result.success is True, f"error: {result.error}"
        assert result.metadata["table_rows"] == 2

        from app.artifacts.manager import get_artifact_manager
        art = result.metadata["artifact"]
        output_path = get_artifact_manager().get_artifact(art["artifact_id"]).physical_path
        from docx import Document
        doc = Document(str(output_path))
        assert len(doc.tables) >= 1, "Expected at least one table in the docx"
        table = doc.tables[0]
        # Header row + 2 data rows = 3 rows
        assert len(table.rows) == 3, f"Expected 3 rows (header + 2 data), got {len(table.rows)}"
        cell_texts = [cell.text for row in table.rows for cell in row.cells]
        assert "F-001" in cell_texts
        assert "F-002" in cell_texts
        assert "HIGH" in cell_texts

    # ── Test 7: missing required args → failure ───────────────────────────

    def test_missing_title_returns_failure(self, tmp_path):
        tool = self._make_tool(tmp_path)
        result = _run(tool.execute(
            sections=[],
            output_filename="out.docx",
        ))
        assert result.success is False
        assert "title" in result.error.lower()


# ===========================================================================
# 8–9: GeneratePptxTool
# ===========================================================================

class TestGeneratePptxTool:

    def _make_tool(self, tmp_path):
        from app.tools.docgen.generate_pptx import GeneratePptxTool
        return GeneratePptxTool(audit_logger=_make_audit(tmp_path))

    # ── Test 8: valid .pptx with correct slide count ──────────────────────

    def test_generates_valid_pptx_with_slides(self, tmp_path):
        """Tool produces a file that python-pptx can open with expected slide count."""
        tool = self._make_tool(tmp_path)

        result = _run(tool.execute(
            title="Inspection Summary",
            slides=[
                {"heading": "Finding F-001", "bullets": ["Temperature anomaly at HX-401", "Severity: HIGH"]},
                {"heading": "Finding F-002", "bullets": ["Pressure variance at OC-12", "Severity: MEDIUM"]},
            ],
            output_filename="test_slides.pptx",
            request_id="pptx-001",
        ))

        assert result.success is True, f"error: {result.error}"
        # result.output is now the sanitized filename; physical path via ArtifactManager
        art = result.metadata["artifact"]
        assert "physical_path" not in art
        assert art["filename"].endswith(".pptx")

        from app.artifacts.manager import get_artifact_manager
        output_path = get_artifact_manager().get_artifact(art["artifact_id"]).physical_path
        assert output_path.exists()
        assert output_path.suffix == ".pptx"

        from pptx import Presentation
        prs = Presentation(str(output_path))
        # Title slide + 2 content slides = 3 total
        assert len(prs.slides) == 3, f"Expected 3 slides, got {len(prs.slides)}"
        assert result.metadata["slide_count"] == 2  # content slides only

    # ── Test 9: slide headings match input ────────────────────────────────

    def test_slide_headings_match_input(self, tmp_path):
        tool = self._make_tool(tmp_path)

        result = _run(tool.execute(
            title="My Presentation",
            slides=[
                {"heading": "Slide Alpha", "bullets": ["Point 1", "Point 2"]},
                {"heading": "Slide Beta", "bullets": ["Point A"]},
            ],
            output_filename="test_headings.pptx",
            request_id="pptx-002",
        ))

        assert result.success is True
        from app.artifacts.manager import get_artifact_manager
        art = result.metadata["artifact"]
        output_path = get_artifact_manager().get_artifact(art["artifact_id"]).physical_path
        from pptx import Presentation
        prs = Presentation(str(output_path))
        titles = [
            sl.shapes.title.text
            for sl in prs.slides
            if sl.shapes.title and sl.shapes.title.text
        ]
        assert "My Presentation" in titles
        assert "Slide Alpha" in titles
        assert "Slide Beta" in titles


# ===========================================================================
# 10–11: GenerateXlsxTool
# ===========================================================================

class TestGenerateXlsxTool:

    def _make_tool(self, tmp_path):
        from app.tools.docgen.generate_xlsx import GenerateXlsxTool
        return GenerateXlsxTool(audit_logger=_make_audit(tmp_path))

    # ── Test 10: valid .xlsx with correct cell values ─────────────────────

    def test_generates_valid_xlsx_with_data(self, tmp_path):
        """Tool produces an .xlsx that openpyxl can open with correct cell values."""
        tool = self._make_tool(tmp_path)

        result = _run(tool.execute(
            sheet_name="Findings",
            headers=["ID", "Finding", "Severity", "Status"],
            rows=[
                ["F-001", "Temperature anomaly", "HIGH", "OPEN"],
                ["F-002", "Pressure variance", "MEDIUM", "OPEN"],
            ],
            output_filename="test_data.xlsx",
            request_id="xlsx-001",
        ))

        assert result.success is True, f"error: {result.error}"
        # result.output is now the sanitized filename; physical path via ArtifactManager
        art = result.metadata["artifact"]
        assert "physical_path" not in art
        assert art["filename"].endswith(".xlsx")

        from app.artifacts.manager import get_artifact_manager
        output_path = get_artifact_manager().get_artifact(art["artifact_id"]).physical_path
        assert output_path.exists()
        assert output_path.suffix == ".xlsx"

        import openpyxl
        wb = openpyxl.load_workbook(str(output_path))
        ws = wb.active
        # Row 1 = headers, Row 2 = first data row
        assert ws["A1"].value == "ID"
        assert ws["C1"].value == "Severity"
        assert ws["A2"].value == "F-001"
        assert ws["C2"].value == "HIGH"
        assert ws["A3"].value == "F-002"
        assert result.metadata["row_count"] == 2
        assert result.metadata["col_count"] == 4

    # ── Test 11: formula row written correctly ─────────────────────────────

    def test_formula_row_written(self, tmp_path):
        """Formula strings in formula_row appear as formulas in the sheet."""
        tool = self._make_tool(tmp_path)

        result = _run(tool.execute(
            headers=["Item", "Value"],
            rows=[
                ["A", 10],
                ["B", 20],
                ["C", 30],
            ],
            formula_row=["Total", "=SUM(B2:B4)"],
            output_filename="test_formula.xlsx",
            request_id="xlsx-002",
        ))

        assert result.success is True
        from app.artifacts.manager import get_artifact_manager
        art = result.metadata["artifact"]
        output_path = get_artifact_manager().get_artifact(art["artifact_id"]).physical_path
        import openpyxl
        wb = openpyxl.load_workbook(str(output_path))
        ws = wb.active
        # Row 5 should be the formula row (header + 3 data + formula)
        assert ws["A5"].value == "Total"
        assert ws["B5"].value == "=SUM(B2:B4)"


# ===========================================================================
# 12–14: Orchestrator integration tests
# ===========================================================================

class TestOrchestratorIntegration:
    """
    Confirm each new Phase D tool is reachable within an orchestrator plan.
    Pattern mirrors the Phase C orchestrator integration test.
    """

    class _FakeResponse:
        content: str
        prompt_tokens: int = 5
        response_tokens: int = 5
        total_tokens: int = 10
        latency_ms: float = 1.0
        model_name: str = "fake"
        ollama_tag: str = "fake:latest"
        request_id: str = "fake-rid"

    class _FakeOllamaClient:
        def __init__(self, responses):
            self._responses = responses
            self._idx = 0

        async def chat_completion(self, messages, *, request_id=None,
                                   temperature=0.0, max_tokens=512,
                                   extra_body=None):
            idx = min(self._idx, len(self._responses) - 1)
            content = self._responses[idx]
            self._idx += 1

            class _R:
                pass

            r = _R()
            r.content = content
            r.prompt_tokens = 5
            r.response_tokens = 5
            r.total_tokens = 10
            r.latency_ms = 1.0
            r.model_name = "fake"
            r.ollama_tag = "fake:latest"
            r.request_id = request_id or "fake"
            return r

    def _make_orch(self, plan_json: str, synthesis: str, tmp_path: Path):
        from app.orchestrator.orchestrator import Orchestrator
        from app.audit.logger import AuditLogger
        audit = AuditLogger(log_path=tmp_path / "audit.jsonl")
        client = self._FakeOllamaClient([plan_json, synthesis])
        return Orchestrator(llm_client=client, audit_logger=audit)

    # ── Test 12: generate_docx callable in a plan ─────────────────────────

    def test_generate_docx_in_plan(self, tmp_path):
        from app.orchestrator.state import OrchestratorStatus

        plan = json.dumps([{
            "step_index": 0,
            "tool_name": "generate_docx",
            "tool_args": {
                "title": "Orchestrator Test Doc",
                "sections": [{"heading": "Sec1", "body": "Body text."}],
                "output_filename": "orch_test.docx",
                "request_id": "orch-docx-001",
            },
            "description": "Generate a test docx",
        }])
        synthesis = "Generated the document successfully."

        orch = self._make_orch(plan, synthesis, tmp_path)
        # _FakeOllamaClient now needs: plan → VERIFIED (sandbox verify call) → synthesis
        orch._llm._responses = [plan, "VERIFIED", synthesis]
        orch._llm._idx = 0
        run = _run(orch.run("Generate a test document", request_id="orch-docx-001"))

        assert run.status == OrchestratorStatus.COMPLETED, (
            f"Expected COMPLETED, got {run.status}. Failure: {run.failure_summary}"
        )
        assert run.outcomes[0].success is True, f"Tool failed: {run.outcomes[0].error}"
        art = run.outcomes[0].metadata.get("artifact")
        assert art is not None, "metadata['artifact'] missing from docx outcome"
        from app.artifacts.manager import get_artifact_manager
        output_path = get_artifact_manager().get_artifact(art["artifact_id"]).physical_path
        assert output_path.exists()


    # ── Test 13: generate_pptx callable in a plan ─────────────────────────

    def test_generate_pptx_in_plan(self, tmp_path):
        from app.orchestrator.state import OrchestratorStatus

        plan = json.dumps([{
            "step_index": 0,
            "tool_name": "generate_pptx",
            "tool_args": {
                "title": "Orchestrator PPTX Test",
                "slides": [{"heading": "Slide 1", "bullets": ["Bullet A"]}],
                "output_filename": "orch_test.pptx",
            },
            "description": "Generate a test pptx",
        }])

        orch = self._make_orch(plan, "Done.", tmp_path)
        run = _run(orch.run("Make a presentation", request_id="orch-pptx-001"))

        assert run.status == OrchestratorStatus.COMPLETED
        assert run.outcomes[0].success is True

    # ── Test 14: code_sandbox callable in a plan ──────────────────────────

    def test_code_sandbox_in_plan(self, tmp_path):
        """
        We don't need Docker for this test — we're only verifying the
        orchestrator can dispatch to the tool.  We mock _run_container
        so Docker is not required.
        """
        from app.orchestrator.state import OrchestratorStatus
        from app.tools.code_sandbox import CodeSandboxTool
        from app.tools.registry import register_tool, TOOL_REGISTRY

        plan = json.dumps([{
            "step_index": 0,
            "tool_name": "code_sandbox",
            "tool_args": {
                "code": "print('hello')",
                "language": "python",
            },
            "description": "Run a trivial script",
        }])
        synthesis = "The script printed hello."

        # Patch the container runner so no real Docker is needed.
        # Must accept `self` — patch.object replaces the unbound class method.
        async def _fake_run(self_inner, container_name, host_script_path):
            return 0, "hello\n", "", False

        # LLM call sequence: plan → VERIFIED (new verify call) → synthesis
        with patch.object(CodeSandboxTool, "_run_container", _fake_run):
            orch = self._make_orch(plan, synthesis, tmp_path)
            orch._llm._responses = [plan, "VERIFIED", synthesis]
            orch._llm._idx = 0
            run = _run(orch.run("Run hello world", request_id="orch-sandbox-001"))

        assert run.status == OrchestratorStatus.COMPLETED
        assert run.outcomes[0].success is True
        output = run.outcomes[0].output
        assert output["stdout"].strip() == "hello"


# ===========================================================================
# 15–18: CodeSandboxTool — real Docker tests
# ===========================================================================
# These tests require a running Docker daemon.
# They are marked with a custom marker so they can be skipped when Docker
# is not available: pytest -m "not docker"
# But per spec, Docker is acceptable as a local dependency (same reasoning
# as local Qdrant in Phase C).

def _docker_available() -> bool:
    """Return True if the Docker daemon is reachable."""
    try:
        r = subprocess.run(
            ["docker", "info"],
            capture_output=True,
            timeout=10,
            check=False,
        )
        return r.returncode == 0
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return False


_DOCKER_SKIP = pytest.mark.skipif(
    not _docker_available(),
    reason="Docker daemon not available — skipping sandbox tests",
)


@_DOCKER_SKIP
class TestCodeSandboxDocker:

    def _make_tool(self, tmp_path):
        from app.tools.code_sandbox import CodeSandboxTool
        return CodeSandboxTool(audit_logger=_make_audit(tmp_path))

    # ── Test 15: trivial print script ─────────────────────────────────────

    def test_trivial_script_succeeds(self, tmp_path):
        tool = self._make_tool(tmp_path)
        result = _run(tool.execute(
            code="print('sovereign-ok')",
            language="python",
            request_id="sandbox-015",
        ))

        assert result.success is True, f"error: {result.error}"
        assert result.output["exit_code"] == 0
        assert "sovereign-ok" in result.output["stdout"]
        assert result.output["timed_out"] is False

    # ── Test 16: exception script → exit 1, stderr captured ───────────────

    def test_exception_script_returns_exit_1(self, tmp_path):
        tool = self._make_tool(tmp_path)
        code = "raise ValueError('deliberate failure')"
        result = _run(tool.execute(
            code=code,
            language="python",
            request_id="sandbox-016",
        ))

        assert result.success is False  # exit_code=1 now correctly propagates as failure
        assert result.output["exit_code"] == 1
        assert "ValueError" in result.output["stderr"]
        assert result.output["timed_out"] is False

    # ── Test 17: sleep-past-timeout script ────────────────────────────────

    def test_timeout_fires(self, tmp_path):
        """A script that sleeps past the timeout must not hang indefinitely."""
        from app.tools import code_sandbox
        original_timeout = code_sandbox._TIMEOUT_SECONDS

        # Use a very short timeout so the test is fast
        code_sandbox._TIMEOUT_SECONDS = 3

        tool = self._make_tool(tmp_path)
        try:
            result = _run(tool.execute(
                code="import time; time.sleep(60)",
                language="python",
                request_id="sandbox-017",
            ))
        finally:
            code_sandbox._TIMEOUT_SECONDS = original_timeout

        assert result.success is False
        assert result.output["timed_out"] is True
        assert "timed out" in result.error.lower()

    # ── Test 18: no leftover containers after each run ────────────────────

    def test_no_leftover_containers(self, tmp_path):
        """
        After running three scripts (success, failure, timeout), confirm
        that no sovereign-sandbox containers remain in `docker ps -a`.
        """
        from app.tools import code_sandbox

        original_timeout = code_sandbox._TIMEOUT_SECONDS
        code_sandbox._TIMEOUT_SECONDS = 3

        tool = self._make_tool(tmp_path)
        try:
            # 1. Success
            _run(tool.execute(code="print('clean')", language="python",
                               request_id="cleanup-1"))
            # 2. Failure
            _run(tool.execute(code="raise RuntimeError('boom')", language="python",
                               request_id="cleanup-2"))
            # 3. Timeout
            _run(tool.execute(code="import time; time.sleep(60)", language="python",
                               request_id="cleanup-3"))
        finally:
            code_sandbox._TIMEOUT_SECONDS = original_timeout

        # Check for leftover containers
        r = subprocess.run(
            ["docker", "ps", "-a", "--filter", "name=sovereign-sandbox",
             "--format", "{{.Names}}"],
            capture_output=True,
            text=True,
            timeout=10,
        )
        leftover = r.stdout.strip()
        assert leftover == "", (
            f"Leftover sandbox containers found after test run:\n{leftover}"
        )
